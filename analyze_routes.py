"""Choose the best depot location(s) and plan delivery routes, writing everything to one Excel file.

Reads the validated CSV produced by validate_data.py, evaluates every combination
of candidate depot sites, plans capacity-constrained delivery routes from each, and
writes a brand-new workbook (no template) with these sheets:
  - Input Data        the validated records the analysis was run on
  - Site Comparison   every candidate site combination ranked by total annual cost
  - Routes            one row per vehicle route: load, distance, duration and stop sequence
  - Route Stops       the stop-by-stop plan for every route
  - Summary           key numbers, including how much the optimized routes save versus a naive plan
  - Route Map         map of the depots, stops and routes (also saved as a PNG)

The optimization is a minimise-cost-under-constraints problem, solved in two nested layers:
  1. Location (facility location): try every combination of --num-sites candidate sites. Each
     customer is served by its nearest chosen depot. The objective is
         annual cost = site fixed costs + operating days x (km driven x cost per km + vehicles x daily vehicle cost)
  2. Routing (capacity-constrained vehicle routing, CVRP), solved separately for each depot:
       - Clarke-Wright savings: start with one route per customer, then repeatedly merge the pair of
         routes with the biggest distance saving, as long as vehicle capacity and the maximum route
         duration are respected.
       - 2-opt: reorder the stops inside each route to remove crossing legs.
Because routing is solved for every candidate combination, the site choice reflects real route costs
rather than a straight-line proxy. Both steps are heuristic / exhaustive-search methods, not a proof of
the global optimum, but they are standard, fast and typically within a few percent of optimal.

Distances are great-circle (haversine) kilometres multiplied by a circuity factor (default 1.3) to
approximate road distance; travel time = distance / average speed + service time at each stop.
"""

import argparse
import csv
import math
import os
from itertools import combinations

import matplotlib

matplotlib.use("Agg")  # no display needed; we only save to file
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from routing_common import (  # noqa: E402
    CSV_FIELDS,
    OUTPUT_MAP_PATH,
    OUTPUT_WORKBOOK_PATH,
    VALIDATED_DATA_PATH,
    Location,
)

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BEST_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

MAX_COMBINATIONS = 2000


class Problem:
    """Everything the solver needs: locations, distance matrix and operating assumptions."""

    def __init__(self, locations: list[Location], args: argparse.Namespace):
        self.locations = locations
        self.customers = [i for i, loc in enumerate(locations) if loc.location_type == "customer"]
        self.sites = [i for i, loc in enumerate(locations) if loc.location_type == "candidate_site"]
        self.capacity = args.vehicle_capacity
        self.max_minutes = args.max_route_hours * 60
        self.speed = args.speed_kmh
        self.cost_per_km = args.cost_per_km
        self.vehicle_daily_cost = args.vehicle_daily_cost
        self.operating_days = args.operating_days
        self.dist = [
            [haversine_km(a, b) * args.circuity for b in locations] for a in locations
        ]

    def demand(self, i: int) -> float:
        return self.locations[i].demand_units

    def route_km(self, depot: int, route: list[int]) -> float:
        path = [depot, *route, depot]
        return sum(self.dist[a][b] for a, b in zip(path, path[1:]))

    def route_minutes(self, depot: int, route: list[int]) -> float:
        service = sum(self.locations[i].service_min for i in route)
        return self.route_km(depot, route) / self.speed * 60 + service


def haversine_km(a: Location, b: Location) -> float:
    r = 6371.0
    p1, p2 = math.radians(a.lat), math.radians(b.lat)
    dphi = p2 - p1
    dlmb = math.radians(b.lon - a.lon)
    h = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlmb / 2) ** 2
    return 2 * r * math.asin(math.sqrt(h))


def read_locations(path: str) -> list[Location]:
    with open(path, newline="", encoding="utf-8") as f:
        return [Location(**row) for row in csv.DictReader(f)]


def clarke_wright(p: Problem, depot: int, customers: list[int]) -> list[list[int]]:
    """Clarke-Wright savings heuristic for a capacity- and duration-constrained set of routes."""
    routes: dict[int, list[int]] = {c: [c] for c in customers}
    route_of = {c: c for c in customers}
    load = {c: p.demand(c) for c in customers}

    savings = sorted(
        ((p.dist[depot][i] + p.dist[depot][j] - p.dist[i][j], i, j) for i, j in combinations(customers, 2)),
        reverse=True,
    )
    for saving, i, j in savings:
        if saving <= 0:
            break
        ri, rj = route_of[i], route_of[j]
        if ri == rj:
            continue
        a, b = routes[ri], routes[rj]
        # Only join routes at their ends: i must end one route and j must start the other (flipping if needed).
        if i not in (a[0], a[-1]) or j not in (b[0], b[-1]):
            continue
        if load[ri] + load[rj] > p.capacity:
            continue
        if a[-1] != i:
            a = a[::-1]
        if b[0] != j:
            b = b[::-1]
        merged = a + b
        if p.route_minutes(depot, merged) > p.max_minutes:
            continue

        routes[ri] = merged
        load[ri] += load.pop(rj)
        del routes[rj]
        for c in b:
            route_of[c] = ri

    return list(routes.values())


def two_opt(p: Problem, depot: int, route: list[int]) -> list[int]:
    """Reverse segments of the route while doing so shortens it."""
    path = [depot, *route, depot]
    improved = True
    while improved:
        improved = False
        for i in range(1, len(path) - 2):
            for j in range(i + 1, len(path) - 1):
                delta = (
                    p.dist[path[i - 1]][path[j]]
                    + p.dist[path[i]][path[j + 1]]
                    - p.dist[path[i - 1]][path[i]]
                    - p.dist[path[j]][path[j + 1]]
                )
                if delta < -1e-9:
                    path[i : j + 1] = reversed(path[i : j + 1])
                    improved = True
    return path[1:-1]


def naive_routes(p: Problem, customers: list[int]) -> list[list[int]]:
    """Baseline: fill vehicles in the order customers appear in the data, with no optimization."""
    routes: list[list[int]] = []
    current: list[int] = []
    for c in customers:
        trial = current + [c]
        if current and (sum(p.demand(i) for i in trial) > p.capacity):
            routes.append(current)
            current = [c]
        else:
            current = trial
    if current:
        routes.append(current)
    return routes


def assign_to_sites(p: Problem, chosen: tuple[int, ...]) -> dict[int, list[int]]:
    assignment: dict[int, list[int]] = {s: [] for s in chosen}
    for c in p.customers:
        assignment[min(chosen, key=lambda s: p.dist[s][c])].append(c)
    return assignment


def plan_routes(p: Problem, chosen: tuple[int, ...], optimize: bool = True) -> list[tuple[int, list[int]]]:
    """Return (depot, route) pairs for the chosen sites."""
    plan = []
    for depot, customers in assign_to_sites(p, chosen).items():
        if not customers:
            continue
        routes = clarke_wright(p, depot, customers) if optimize else naive_routes(p, customers)
        for r in routes:
            plan.append((depot, two_opt(p, depot, r) if optimize else r))
    return plan


def plan_cost(p: Problem, chosen: tuple[int, ...], plan: list[tuple[int, list[int]]]) -> dict:
    km = sum(p.route_km(d, r) for d, r in plan)
    vehicles = len(plan)
    daily = km * p.cost_per_km + vehicles * p.vehicle_daily_cost
    fixed = sum(p.locations[s].fixed_annual_cost for s in chosen)
    return {
        "daily_km": km,
        "vehicles": vehicles,
        "annual_transport_cost": daily * p.operating_days,
        "fixed_site_cost": fixed,
        "annual_total_cost": daily * p.operating_days + fixed,
    }


def evaluate_sites(p: Problem, num_sites: int) -> list[dict]:
    results = []
    for chosen in combinations(p.sites, num_sites):
        plan = plan_routes(p, chosen)
        results.append({"chosen": chosen, "plan": plan, **plan_cost(p, chosen, plan)})
    results.sort(key=lambda r: r["annual_total_cost"])
    for rank, r in enumerate(results, start=1):
        r["rank"] = rank
    return results


def build_site_comparison(p: Problem, results: list[dict]) -> list[dict]:
    best_cost = results[0]["annual_total_cost"]
    return [
        {
            "rank": r["rank"],
            "sites": " + ".join(p.locations[s].name for s in r["chosen"]),
            "site_ids": " + ".join(p.locations[s].location_id for s in r["chosen"]),
            "fixed_site_cost": round(r["fixed_site_cost"]),
            "daily_km": round(r["daily_km"], 1),
            "vehicles": r["vehicles"],
            "annual_transport_cost": round(r["annual_transport_cost"]),
            "annual_total_cost": round(r["annual_total_cost"]),
            "vs_best": round(r["annual_total_cost"] - best_cost),
        }
        for r in results
    ]


def build_route_tables(p: Problem, plan: list[tuple[int, list[int]]]) -> tuple[list[dict], list[dict]]:
    routes, stops = [], []
    ordered = sorted(plan, key=lambda dr: (dr[0], p.locations[dr[1][0]].location_id))
    for n, (depot, route) in enumerate(ordered, start=1):
        load = sum(p.demand(i) for i in route)
        minutes = p.route_minutes(depot, route)
        depot_loc = p.locations[depot]
        routes.append(
            {
                "route_id": f"R{n:02d}",
                "depot": depot_loc.name,
                "stops": len(route),
                "load_units": load,
                "load_pct": round(load / p.capacity * 100, 1),
                "distance_km": round(p.route_km(depot, route), 1),
                "duration_hours": round(minutes / 60, 2),
                "sequence": " > ".join(["DEPOT", *(p.locations[i].location_id for i in route), "DEPOT"]),
            }
        )
        cumulative = 0.0
        for seq, i in enumerate(route, start=1):
            loc = p.locations[i]
            cumulative += loc.demand_units
            stops.append(
                {
                    "route_id": f"R{n:02d}",
                    "seq": seq,
                    "location_id": loc.location_id,
                    "name": loc.name,
                    "lat": loc.lat,
                    "lon": loc.lon,
                    "demand_units": loc.demand_units,
                    "cumulative_load": cumulative,
                }
            )
    return routes, stops


def build_summary(p: Problem, results: list[dict], routes: list[dict], num_sites: int) -> list[dict]:
    best = results[0]
    naive_plan = plan_routes(p, best["chosen"], optimize=False)
    naive = plan_cost(p, best["chosen"], naive_plan)
    km_saving = 1 - best["daily_km"] / naive["daily_km"] if naive["daily_km"] else 0.0
    cost_saving = naive["annual_total_cost"] - best["annual_total_cost"]
    total_demand = sum(p.demand(i) for i in p.customers)

    rows = [
        {"metric": f"Selected site(s) (best of {len(results)} combinations)", "value": " + ".join(p.locations[s].name for s in best["chosen"])},
        {"metric": "Delivery stops", "value": len(p.customers)},
        {"metric": "Total daily demand (units)", "value": round(total_demand)},
        {"metric": "Vehicles needed", "value": best["vehicles"]},
        {"metric": "Vehicle capacity (units)", "value": p.capacity},
        {"metric": "Average vehicle load", "value": f"{sum(r['load_pct'] for r in routes) / len(routes):.0f}%"},
        {"metric": "Longest route (hours)", "value": max(r["duration_hours"] for r in routes)},
        {"metric": "Optimized distance per day (km)", "value": round(best["daily_km"], 1)},
        {"metric": "Naive plan distance per day (km)", "value": round(naive["daily_km"], 1)},
        {"metric": "Distance saved by route optimization", "value": f"{km_saving:.0%}"},
        {"metric": "Naive plan vehicles", "value": naive["vehicles"]},
        {"metric": "Annual transport cost (optimized)", "value": round(best["annual_transport_cost"])},
        {"metric": "Annual site fixed cost", "value": round(best["fixed_site_cost"])},
        {"metric": "Annual total cost (optimized)", "value": round(best["annual_total_cost"])},
        {"metric": "Annual total cost (naive routes, same site)", "value": round(naive["annual_total_cost"])},
        {"metric": "Annual saving from route optimization", "value": round(cost_saving)},
    ]
    if len(results) > 1:
        worst = results[-1]
        rows.append({"metric": "Annual cost of the worst site choice", "value": round(worst["annual_total_cost"])})
        rows.append({"metric": "Annual saving of best vs. worst site choice", "value": round(worst["annual_total_cost"] - best["annual_total_cost"])})
        rows.append({"metric": "Runner-up site(s)", "value": " + ".join(p.locations[s].name for s in results[1]["chosen"])})
        rows.append({"metric": "Runner-up extra annual cost", "value": round(results[1]["annual_total_cost"] - best["annual_total_cost"])})
    return rows


def draw_map(p: Problem, best: dict, path: str) -> None:
    fig, ax = plt.subplots(figsize=(9.5, 8))
    colors = plt.get_cmap("tab20").colors

    ordered = sorted(best["plan"], key=lambda dr: (dr[0], p.locations[dr[1][0]].location_id))
    for n, (depot, route) in enumerate(ordered):
        pts = [p.locations[i] for i in [depot, *route, depot]]
        color = colors[n % len(colors)]
        ax.plot([q.lon for q in pts], [q.lat for q in pts], "-", color=color, linewidth=1.8, alpha=0.85, zorder=2)
        ax.scatter(
            [p.locations[i].lon for i in route],
            [p.locations[i].lat for i in route],
            s=[20 + 6 * p.demand(i) for i in route],
            color=color,
            edgecolor="white",
            linewidth=0.8,
            zorder=3,
        )

    unused = [s for s in p.sites if s not in best["chosen"]]
    ax.scatter([p.locations[s].lon for s in unused], [p.locations[s].lat for s in unused], marker="^", s=80, color="#9e9e9e", edgecolor="white", zorder=3, label="Candidate site (not chosen)")
    for s in unused:
        ax.annotate(p.locations[s].location_id, (p.locations[s].lon, p.locations[s].lat), textcoords="offset points", xytext=(6, 5), fontsize=8, color="#616161")
    ax.scatter([p.locations[s].lon for s in best["chosen"]], [p.locations[s].lat for s in best["chosen"]], marker="*", s=420, color="#d9534f", edgecolor="black", linewidth=0.8, zorder=5, label="Selected depot")

    mean_lat = sum(loc.lat for loc in p.locations) / len(p.locations)
    ax.set_aspect(1 / math.cos(math.radians(mean_lat)))
    ax.set_xlabel("Longitude")
    ax.set_ylabel("Latitude")
    ax.set_title(f"{best['vehicles']} routes, {best['daily_km']:.0f} km/day  (stop size = demand, one colour per route)")
    ax.grid(alpha=0.25)
    ax.legend(loc="lower right", fontsize=8)
    fig.tight_layout()
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    fig.savefig(path, dpi=130)
    plt.close(fig)


def autosize_columns(ws) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, 70)


def write_sheet(ws, fieldnames: list[str], rows: list[dict]) -> None:
    ws.append(fieldnames)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(f, "") for f in fieldnames])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(fieldnames))}{len(rows) + 1}"
    autosize_columns(ws)


def build_workbook(input_rows, site_comparison, routes, stops, summary, map_path) -> Workbook:
    wb = Workbook()

    ws_input = wb.active
    ws_input.title = "Input Data"
    write_sheet(ws_input, CSV_FIELDS, input_rows)

    ws_sites = wb.create_sheet("Site Comparison")
    site_fields = ["rank", "sites", "site_ids", "fixed_site_cost", "daily_km", "vehicles", "annual_transport_cost", "annual_total_cost", "vs_best"]
    write_sheet(ws_sites, site_fields, site_comparison)
    for cell in ws_sites[2]:
        cell.fill = BEST_FILL

    ws_routes = wb.create_sheet("Routes")
    write_sheet(ws_routes, ["route_id", "depot", "stops", "load_units", "load_pct", "distance_km", "duration_hours", "sequence"], routes)

    ws_stops = wb.create_sheet("Route Stops")
    write_sheet(ws_stops, ["route_id", "seq", "location_id", "name", "lat", "lon", "demand_units", "cumulative_load"], stops)

    ws_summary = wb.create_sheet("Summary")
    write_sheet(ws_summary, ["metric", "value"], summary)

    ws_map = wb.create_sheet("Route Map")
    ws_map.add_image(XLImage(map_path), "A1")

    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description="Pick the best depot site(s) and plan delivery routes; write results to Excel.")
    parser.add_argument("--input", type=str, default=VALIDATED_DATA_PATH, help="Path to the validated CSV")
    parser.add_argument("--output", type=str, default=OUTPUT_WORKBOOK_PATH, help="Path to write the Excel workbook to")
    parser.add_argument("--map", type=str, default=OUTPUT_MAP_PATH, help="Path to write the route map PNG to")
    parser.add_argument("--num-sites", type=int, default=1, help="How many depots to open (default: 1)")
    parser.add_argument("--vehicle-capacity", type=float, default=100, help="Vehicle capacity in units (default: 100)")
    parser.add_argument("--max-route-hours", type=float, default=8, help="Maximum driving + service time per route (default: 8)")
    parser.add_argument("--speed-kmh", type=float, default=30, help="Average driving speed in km/h (default: 30)")
    parser.add_argument("--circuity", type=float, default=1.3, help="Road distance / straight-line distance factor (default: 1.3)")
    parser.add_argument("--cost-per-km", type=float, default=0.65, help="Running cost per km driven (default: 0.65)")
    parser.add_argument("--vehicle-daily-cost", type=float, default=120, help="Daily fixed cost of each vehicle and driver (default: 120)")
    parser.add_argument("--operating-days", type=int, default=250, help="Delivery days per year (default: 250)")
    args = parser.parse_args()

    locations = read_locations(args.input)
    with open(args.input, newline="", encoding="utf-8") as f:
        input_rows = list(csv.DictReader(f))

    p = Problem(locations, args)
    if not p.customers or not p.sites:
        raise SystemExit("Need at least one customer and one candidate_site in the data.")
    if not 1 <= args.num_sites <= len(p.sites):
        raise SystemExit(f"--num-sites must be between 1 and the number of candidate sites ({len(p.sites)}).")
    too_big = [p.locations[i].location_id for i in p.customers if p.demand(i) > p.capacity]
    if too_big:
        raise SystemExit(f"Demand exceeds vehicle capacity ({p.capacity}) for: {', '.join(too_big)}")
    n_combos = math.comb(len(p.sites), args.num_sites)
    if n_combos > MAX_COMBINATIONS:
        raise SystemExit(f"{n_combos} site combinations is too many to evaluate exhaustively (limit {MAX_COMBINATIONS}); reduce --num-sites or the candidate list.")

    results = evaluate_sites(p, args.num_sites)
    best = results[0]
    routes, stops = build_route_tables(p, best["plan"])
    summary = build_summary(p, results, routes, args.num_sites)
    site_comparison = build_site_comparison(p, results)

    draw_map(p, best, args.map)
    wb = build_workbook(input_rows, site_comparison, routes, stops, summary, args.map)
    os.makedirs(os.path.dirname(args.output) or ".", exist_ok=True)
    wb.save(args.output)

    names = " + ".join(p.locations[s].name for s in best["chosen"])
    print(f"Evaluated {len(results)} site combination(s) for {len(p.customers)} stops -> {args.output}")
    print(f"Best: {names} ({best['vehicles']} vehicles, {best['daily_km']:.0f} km/day, {best['annual_total_cost']:,.0f} total annual cost)")
    long_routes = [r for r in routes if r["duration_hours"] > args.max_route_hours + 1e-9]
    if long_routes:
        print(f"Warning: {len(long_routes)} route(s) exceed the {args.max_route_hours}h limit (single stops too far for the limit).")


if __name__ == "__main__":
    main()
